#!/usr/bin/env python3
"""Exports tests/cases/test-cases.xlsx to docs/test-cases.md (for PR review) and checks coverage:
every REQ in docs/product-brief.md needs a Positive, a Negative and an Edge case; every NFR needs a row.
Usage: python tests/tools/export_cases.py [--check]   (exit 1 when coverage has gaps)"""
import re, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SHEET = ROOT / 'tests/cases/test-cases.xlsx'
OUT = ROOT / 'docs/test-cases.md'
FUNCTIONAL = ['Web', 'Mobile', 'API', 'Payments']


def cell(v):
    return '' if v is None else str(v).replace('|', '\\|').replace('\n', ' ')


def main():
    wb = openpyxl.load_workbook(SHEET, data_only=False)
    brief = (ROOT / 'docs/product-brief.md').read_text(encoding='utf-8')
    reqs = sorted(set(re.findall(r'^\| (REQ-\d+) \|', brief, re.M)))
    nfrs = sorted(set(re.findall(r'^\| (NFR-\d+) \|', brief, re.M)))
    types = {r: set() for r in reqs}
    nfr_rows = set()
    lines = ['# Test cases', '', f'Generated from `tests/cases/test-cases.xlsx` by `tests/tools/export_cases.py`. One row per step.', '']
    totals = []
    for name in FUNCTIONAL + ['Non-functional']:
        ws = wb[name]
        head = [c.value for c in ws[1]]
        rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        cases = [r for r in rows if r[8] == 1]
        auto = sum(1 for r in cases if (r[13] if name == 'Non-functional' else r[12]) == 'Y')
        totals.append((name, len(cases), auto))
        lines += [f'## {name} ({len(cases)} cases, {auto} automated)', '', '| ' + ' | '.join(cell(h) for h in head) + ' |', '|' + '---|' * len(head)]
        lines += ['| ' + ' | '.join(cell(v) for v in r[:len(head)]) + ' |' for r in rows]
        lines.append('')
        for r in cases:
            if name == 'Non-functional':
                nfr_rows.update(re.findall(r'NFR-\d+', cell(r[2])))
                continue
            for q in re.findall(r'REQ-\d+', cell(r[2])):
                types.setdefault(q, set()).add(r[4])
            for a, b in re.findall(r'REQ-(\d+)–REQ-(\d+)', cell(r[2])):
                for k in range(int(a), int(b) + 1):
                    types.setdefault(f'REQ-{k:03d}', set()).add(r[4])
    gaps = {q: sorted({'Positive', 'Negative', 'Edge'} - t) for q, t in types.items() if q in reqs and {'Positive', 'Negative', 'Edge'} - t}
    missing_nfr = [n for n in nfrs if n not in nfr_rows]
    summary = ['## Summary', '', '| Sheet | Cases | Automated |', '|---|---|---|'] + [f'| {n} | {c} | {a} |' for n, c, a in totals]
    summary += ['', f'REQs: {len(reqs)} · with Positive, Negative and Edge: {len(reqs) - len(gaps)} · NFRs with rows: {len(nfrs) - len(missing_nfr)}/{len(nfrs)}', '']
    OUT.write_text('\n'.join(lines[:4] + summary + lines[4:]) + '\n', encoding='utf-8')
    print(f'wrote {OUT.relative_to(ROOT)}; ' + ', '.join(f'{n} {c}' for n, c, _ in totals))
    if gaps or missing_nfr:
        print('coverage gaps:', gaps, missing_nfr)
        if '--check' in sys.argv:
            sys.exit(1)
    else:
        print('coverage complete: every REQ has Positive, Negative and Edge; every NFR has a row')


if __name__ == '__main__':
    main()
