#!/usr/bin/env python3
"""Bug log and test-case status updater (needs: pip install openpyxl).

  python tests/tools/bug_log.py init
      Create tests/bugs/bug-log.xlsx if it does not exist.
  python tests/tools/bug_log.py ingest test-results/summary.json [--report docs/test-reports/x.md]
      Add a Runs row; log every failed/flaky result as a bug (new, updated or reopened);
      mark Fixed bugs whose test now passes on staging as Verified; write Pass/Fail/Flaky
      into the Status column of tests/cases/test-cases.xlsx. Prints a JSON digest.
  python tests/tools/bug_log.py list [open|all]
      Print bugs as JSON.
  python tests/tools/bug_log.py set BUG-0003 "Status=In PR" "Fix branch / PR=https://..."
      Update fields of one bug (column names exactly as in the header row).
"""
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
BUGS = ROOT / "tests" / "bugs" / "bug-log.xlsx"
CASES = ROOT / "tests" / "cases" / "test-cases.xlsx"

BUG_COLS = [
    ("Bug ID", 11), ("Status", 12), ("Decision", 12), ("Severity", 9), ("Priority", 9),
    ("Environment", 11), ("Platform", 22), ("Suite", 12), ("Kind", 14), ("TC ID", 11),
    ("REQ / NFR ID", 12), ("Title", 40), ("Steps to reproduce", 50), ("Expected result", 40),
    ("Actual result", 50), ("Evidence", 34), ("First seen", 11), ("First run", 18),
    ("Last seen", 11), ("Last run", 18), ("Occurrences", 11), ("Fix branch / PR", 34),
    ("Fixed in", 16), ("Verified in run", 18), ("Signature", 30), ("Remarks", 30),
]
RUN_COLS = [
    ("Run ID", 20), ("Date", 11), ("Environment", 11), ("Base URL", 32), ("Commit", 10),
    ("Run link", 40), ("Total", 8), ("Passed", 8), ("Failed", 8), ("Flaky", 8), ("Warn", 8),
    ("Skipped", 8), ("New bugs", 9), ("Reopened", 9), ("Still open", 9), ("Verified", 9),
    ("Report", 40),
]
STATUSES = ["New", "Open", "In progress", "In PR", "Fixed", "Verified", "Closed",
            "Reopened", "Deferred", "Won't fix", "Not a bug", "Duplicate"]
OPEN = {"New", "Open", "In progress", "In PR", "Reopened", "Deferred"}
FIXED = {"Fixed", "Verified", "Closed"}
DECISIONS = ["Fix now", "Defer", "Won't fix", "Not a bug"]
SEVERITIES = ["S1", "S2", "S3", "S4"]
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")


def style_header(ws, cols):
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(cols)).column_letter}1"


def add_list(ws, values, col_letter):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}5000")


def init():
    if BUGS.exists():
        return
    BUGS.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    rd = wb.active
    rd.title = "Read me"
    for line in [
        "Equisteg bug log — one row per bug. Filled by /test after every staging run; edited by /change, /update and /test fixes.",
        "Status flow: New -> Open (Decision = Fix now) -> In progress -> In PR -> Fixed (merged to development) -> Verified (passes on staging) -> Closed.",
        "Other statuses: Deferred (decided to fix later), Won't fix, Not a bug (test or data problem), Duplicate (see Remarks), Reopened (failed again after Fixed/Verified).",
        "Severity: S1 = payments, security, data loss or app unusable; S2 = main feature broken, no workaround; S3 = feature partly broken or workaround exists; S4 = cosmetic, flaky test, minor.",
        "Signature = suite | TC ID or automation ref. The same failure in a later run updates the same row (Last seen, Occurrences) instead of adding a new one.",
        "Never delete rows. Never paste secrets, real customer data or production tokens into this sheet.",
    ]:
        rd.append([line])
    rd.column_dimensions["A"].width = 140

    ws = wb.create_sheet("Bugs")
    style_header(ws, BUG_COLS)
    add_list(ws, STATUSES, "B")
    add_list(ws, DECISIONS, "C")
    add_list(ws, SEVERITIES, "D")
    add_list(ws, ["P1", "P2", "P3"], "E")
    add_list(ws, ["Functional", "Non-functional"], "I")

    runs = wb.create_sheet("Runs")
    style_header(runs, RUN_COLS)

    sm = wb.create_sheet("Summary")
    style_header(sm, [("Status", 14), ("Count", 10), ("", 4), ("Severity (open)", 16), ("Count", 10)])
    for i, s in enumerate(STATUSES, 2):
        sm.cell(row=i, column=1, value=s)
        sm.cell(row=i, column=2, value=f"=COUNTIFS(Bugs!B:B,\"{s}\")")
    sm.cell(row=len(STATUSES) + 2, column=1, value="Total").font = Font(bold=True)
    sm.cell(row=len(STATUSES) + 2, column=2, value=f"=SUM(B2:B{len(STATUSES) + 1})").font = Font(bold=True)
    open_list = [s for s in STATUSES if s in OPEN]
    for i, sev in enumerate(SEVERITIES, 2):
        sm.cell(row=i, column=4, value=sev)
        sm.cell(row=i, column=5, value="=" + "+".join(f"COUNTIFS(Bugs!D:D,\"{sev}\",Bugs!B:B,\"{s}\")" for s in open_list))
    wb.save(BUGS)
    print(f"created {BUGS.relative_to(ROOT)}")


def header_index(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value}


def read_rows(ws):
    idx = header_index(ws)
    for r in range(2, ws.max_row + 1):
        row = {k: ws.cell(row=r, column=c).value for k, c in idx.items()}
        if any(v not in (None, "") for v in row.values()):
            yield r, row


def load_cases():
    """TC ID -> {sheet, row, req, steps, expected, priority}; plus automation-ref -> TC ID."""
    cases, refs = {}, []
    if not CASES.exists():
        return None, cases, refs
    wb = load_workbook(CASES)
    for ws in wb.worksheets:
        idx = header_index(ws)
        if "TC ID" not in idx:
            continue
        req_col = "REQ ID" if "REQ ID" in idx else "NFR ID"
        for r, row in read_rows(ws):
            tc = row.get("TC ID")
            if not tc:
                continue
            c = cases.setdefault(tc, {"sheet": ws.title, "row": r, "req": row.get(req_col),
                                      "priority": row.get("Priority"), "steps": [], "expected": []})
            if row.get("Test step"):
                c["steps"].append(f"{row.get('Step No') or len(c['steps']) + 1}. {row['Test step']}")
            if row.get("Expected result"):
                c["expected"].append(f"{row.get('Step No') or len(c['expected']) + 1}. {row['Expected result']}")
            if row.get("Automation ref"):
                refs.append((str(row["Automation ref"]).strip(), tc))
    return wb, cases, refs


def match_tc(result, refs):
    if result.get("tcId"):
        return result["tcId"]
    ref = (result.get("ref") or "").strip()
    for sheet_ref, tc in refs:
        if sheet_ref == ref:
            return tc
    file_part = ref.split(" :: ")[0]
    name_part = ref.split(" :: ")[1] if " :: " in ref else ""
    for sheet_ref, tc in refs:
        s_file, _, s_name = sheet_ref.partition(" :: ")
        if s_file.strip() == file_part and (not s_name or s_name.strip() in name_part):
            return tc
    return None


def guess_severity(result, case):
    text = f"{result.get('suite')} {result.get('title')} {result.get('ref')}".lower()
    if result["status"] == "flaky":
        return "S4"
    if re.search(r"payment|webhook|entitle|licen|zap high|security|secret|auth", text):
        return "S1"
    pr = (case or {}).get("priority")
    return {"P1": "S2", "P2": "S3", "P3": "S4"}.get(pr, "S3")


def next_bug_id(ws):
    nums = [int(m.group(1)) for _, row in read_rows(ws) if (m := re.match(r"BUG-(\d+)", str(row.get("Bug ID") or "")))]
    return f"BUG-{(max(nums) + 1 if nums else 1):04d}"


def ingest(summary_path, report=""):
    init()
    s = json.loads(Path(summary_path).read_text(encoding="utf-8"))
    run_id, env, today = s["runId"], s.get("env", "staging"), date.today().isoformat()
    cases_wb, cases, refs = load_cases()
    wb = load_workbook(BUGS)
    ws = wb["Bugs"]
    col = header_index(ws)
    by_sig = {row["Signature"]: r for r, row in read_rows(ws) if row.get("Signature")}
    digest = {"runId": run_id, "env": env, "new": [], "reopened": [], "updated": [], "verified": []}

    # Group results by signature so one failure across several browsers/devices is one bug.
    groups = {}
    for res in s.get("results", []):
        tc = match_tc(res, refs)
        res["tcId"] = tc
        sig = f"{res['suite']} | {tc or res.get('ref') or res.get('title')}"
        groups.setdefault(sig, []).append(res)

    tc_status = {}
    for sig, rs in groups.items():
        failing = [r for r in rs if r["status"] in ("fail", "flaky")]
        tc = rs[0]["tcId"]
        if tc:
            st = "Fail" if any(r["status"] == "fail" for r in rs) else "Flaky" if failing else \
                 "Pass" if any(r["status"] == "pass" for r in rs) else None
            if st:
                tc_status[tc] = "Fail" if "Fail" in (tc_status.get(tc), st) else st
        row = by_sig.get(sig)
        get = (lambda k: ws.cell(row=row, column=col[k]).value) if row else None
        put = (lambda k, v: ws.cell(row=row, column=col[k], value=v)) if row else None

        if not failing:
            if row and env == "staging" and get("Status") in ("Fixed", "In PR") and all(r["status"] == "pass" for r in rs):
                put("Status", "Verified")
                put("Verified in run", run_id)
                digest["verified"].append(get("Bug ID"))
            continue

        r0 = failing[0]
        platforms = sorted({r["platform"] for r in failing})
        actual = "\n".join(f"[{r['platform']}] {r['error']}" for r in failing)[:1500]
        evidence = ", ".join(filter(None, {r.get("evidence") for r in failing}))[:500]
        if row:
            put("Last seen", today)
            put("Last run", run_id)
            put("Occurrences", (get("Occurrences") or 0) + 1)
            put("Actual result", actual)
            put("Platform", ", ".join(sorted(set(str(get("Platform") or "").split(", ")) - {""} | set(platforms))))
            if evidence:
                put("Evidence", evidence)
            if get("Status") in FIXED:
                put("Status", "Reopened")
                put("Remarks", f"{get('Remarks') or ''} Reopened by run {run_id}.".strip())
                digest["reopened"].append(get("Bug ID"))
            else:
                digest["updated"].append(get("Bug ID"))
            continue

        case = cases.get(tc) if tc else None
        bug_id = next_bug_id(ws)
        new = {
            "Bug ID": bug_id, "Status": "New", "Severity": guess_severity(r0, case),
            "Priority": (case or {}).get("priority") or "P2", "Environment": env,
            "Platform": ", ".join(platforms), "Suite": r0["suite"],
            "Kind": "Non-functional" if r0["kind"] == "non-functional" or (case or {}).get("sheet") == "Non-functional" else "Functional",
            "TC ID": tc, "REQ / NFR ID": (case or {}).get("req"),
            "Title": ("Flaky: " if r0["status"] == "flaky" else "") + r0["title"][:200],
            "Steps to reproduce": "\n".join((case or {}).get("steps") or [f"Run {r0.get('ref')} against {s.get('baseUrl') or env}"]),
            "Expected result": "\n".join((case or {}).get("expected") or ["Test passes"]),
            "Actual result": actual, "Evidence": evidence or s.get("runUrl"),
            "First seen": today, "First run": run_id, "Last seen": today, "Last run": run_id,
            "Occurrences": 1, "Signature": sig,
        }
        r = ws.max_row + 1
        for k, v in new.items():
            ws.cell(row=r, column=col[k], value=v).alignment = Alignment(wrap_text=True, vertical="top")
        by_sig[sig] = r
        digest["new"].append(bug_id)

    still_open = [row for _, row in read_rows(ws) if row.get("Status") in OPEN]
    t = s.get("totals", {})
    wb["Runs"].append([run_id, today, env, s.get("baseUrl"), (s.get("commit") or "")[:7], s.get("runUrl"),
                       t.get("total"), t.get("passed"), t.get("failed"), t.get("flaky"), t.get("warn"),
                       t.get("skipped"), len(digest["new"]), len(digest["reopened"]), len(still_open),
                       len(digest["verified"]), report])
    wb.save(BUGS)

    if cases_wb and tc_status:
        for tc, st in tc_status.items():
            c = cases.get(tc)
            if not c:
                continue
            sheet = cases_wb[c["sheet"]]
            idx = header_index(sheet)
            sheet.cell(row=c["row"], column=idx["Status"], value=st)
            if "Remarks" in idx:
                sheet.cell(row=c["row"], column=idx["Remarks"], value=f"{env} run {run_id} ({today})")
        cases_wb.save(CASES)

    digest["open"] = [{k: row.get(k) for k in ("Bug ID", "Status", "Severity", "Priority", "Platform", "Kind", "TC ID", "Title", "Decision")}
                      for row in still_open]
    digest["unmappedFailures"] = [r["title"] for rs in groups.values() for r in rs if r["status"] == "fail" and not r["tcId"]]
    print(json.dumps(digest, indent=2, default=str))


def list_bugs(which="open"):
    init()
    ws = load_workbook(BUGS)["Bugs"]
    rows = [row for _, row in read_rows(ws) if which == "all" or row.get("Status") in OPEN]
    print(json.dumps(rows, indent=2, default=str))


def set_fields(bug_id, pairs):
    wb = load_workbook(BUGS)
    ws = wb["Bugs"]
    col = header_index(ws)
    for r, row in read_rows(ws):
        if row.get("Bug ID") == bug_id:
            for p in pairs:
                k, _, v = p.partition("=")
                if k not in col:
                    sys.exit(f"unknown column: {k}")
                ws.cell(row=r, column=col[k], value=v)
            wb.save(BUGS)
            print(f"{bug_id} updated")
            return
    sys.exit(f"{bug_id} not found")


if __name__ == "__main__":
    a = sys.argv[1:]
    if not a or a[0] not in ("init", "ingest", "list", "set"):
        sys.exit(__doc__)
    if a[0] == "init":
        init()
    elif a[0] == "ingest":
        rep = a[a.index("--report") + 1] if "--report" in a else ""
        ingest(a[1], rep)
    elif a[0] == "list":
        list_bugs(a[1] if len(a) > 1 else "open")
    else:
        set_fields(a[1], a[2:])
